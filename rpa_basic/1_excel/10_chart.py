from openpyxl import load_workbook
wb= load_workbook('sample.xlsx')
ws = wb.active
from openpyxl.chart import BarChart, Reference

bar_value = Reference(ws, min_row =2, max_row=11, min_col =2, max_col =3)
bar_chart = BarChart()
bar_chart.add_data(bar_value)

ws.add_chart(bar_chart, "E1")

wb.save("sample_chart.xlsx")